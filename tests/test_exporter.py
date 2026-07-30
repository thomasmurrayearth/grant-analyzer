"""The XLSX is the artefact users forward to co-founders and boards, so it has
to carry the analysis *and* the offer of help with it."""

import io
import unittest

from openpyxl import load_workbook

from exporter import generate_xlsx

RESULT = {
    "company_profile": {"name": "Example Thermal"},
    "executive_summary": {
        "company_overview": "A climate hardware company.",
        "key_constraints": "Pre-revenue, small team.",
    },
    "strategic_recommendations": {
        "best_fit_strategy": "Lead with national innovation grants.",
        "key_risks": "Consortium-only calls need a partner.",
    },
    "opportunities": [{
        "name": "Example Innovation Fund",
        "managing_body": "Example Agency",
        "geography": "EU",
        "thematic_fit_score": 5,
        "strategic_value_score": 4,
        "ease_score": 3,
        "priority_score": 4.5,
        "priority_tier": "Must Pursue",
        "application_link": "https://example.org/apply",
        "deadline": "2026-09-30",
    }],
    "strategic_watchlist": [{
        "name": "Example Relationship Fund",
        "managing_body": "Example Foundation",
        "geography": "UK",
        "opportunity_type": "relationship_led_funder",
        "application_route": "invitation_or_relationship_led",
        "application_timing": "recurring_uncertain",
        "status": "Recurring",
        "application_link": "https://example.org/watch",
        "funding_type": "Grant",
        "max_funding": "£50,000",
        "thematic_relevance": "Strong alignment with the circular economy theme.",
        "why_watchlist": "Requires an existing relationship with the funder.",
        "what_would_unlock": "An introduction from an existing grantee.",
    }, {
        "name": "Example Discovery-Only Programme",
        "managing_body": "Example Council",
        "geography": "NZ",
        "application_link": "https://example.org/discovery-only",
        "funding_type": "Grant",
        "max_funding": "£20,000",
        # Never deep-researched, so no scored thematic_fit_score — only the
        # discovery-stage rating the analyzer backfills onto "thematic_fit".
        "thematic_fit": 4,
        "thematic_relevance": "Discovered but not selected for deep research.",
        "why_watchlist": "Not selected for deep research this run.",
        "what_would_unlock": "A future run with more shortlist slots available.",
    }],
    "acronym_definitions": [{"acronym": "TRL", "definition": "Technology Readiness Level"}],
}


class ExporterTest(unittest.TestCase):
    def setUp(self):
        self.wb = load_workbook(io.BytesIO(generate_xlsx(RESULT)))

    def test_workbook_has_the_three_tabs(self):
        self.assertEqual(
            self.wb.sheetnames,
            ["Grant Opportunities", "Strategic Recommendations", "Definitions acronyms"],
        )

    def test_recommendations_tab_carries_the_analysis(self):
        text = self._sheet_text("Strategic Recommendations")
        self.assertIn("Lead with national innovation grants.", text)
        self.assertIn("A climate hardware company.", text)

    def test_recommendations_tab_carries_the_consulting_offer(self):
        text = self._sheet_text("Strategic Recommendations")
        self.assertIn("Thomas Murray", text)
        self.assertIn("thomasmurray.earth", text)

    def test_opportunities_tab_still_scores_grants(self):
        ws = self.wb["Grant Opportunities"]
        self.assertEqual(ws["A2"].value, "Example Innovation Fund")
        self.assertEqual(ws["B2"].value, "=(2*D2)+F2+H2")

    def test_watchlist_items_are_folded_into_the_same_tab(self):
        # No separate "Strategic Watchlist" tab — still exactly three sheets.
        self.assertEqual(len(self.wb.sheetnames), 3)

        ws = self.wb["Grant Opportunities"]
        # Scored shortlist row sorts above the unscored watchlist row.
        self.assertEqual(ws["A2"].value, "Example Innovation Fund")
        self.assertEqual(ws["A3"].value, "Example Relationship Fund")

        # Watchlist row has no thematic fit rating at all (scored or
        # discovery-stage), so Priority Score reads "unknown" rather than a
        # formula or number — no tier column needed.
        self.assertEqual(ws["B3"].value, "unknown")
        self.assertEqual(ws["D3"].value, "unknown")
        # Strategic value / ease of execution default to 0, not "unknown",
        # even when thematic fit itself is unknown.
        self.assertEqual(ws["F3"].value, 0)
        self.assertEqual(ws["H3"].value, 0)

        self.assertEqual(ws["K3"].value, "https://example.org/watch")
        self.assertEqual(ws["L3"].value, "Example Foundation")
        self.assertEqual(ws["T3"].value, "UK")
        self.assertEqual(ws["X3"].value, "Recurring")
        self.assertIn("circular economy", ws["C3"].value)
        self.assertIn("relationship with the funder", ws["V3"].value)

    def test_watchlist_item_with_discovery_stage_fit_gets_a_priority_score(self):
        # A watchlist item carrying only the discovery-stage "thematic_fit"
        # (no scored thematic_fit_score) still gets a real Thematic Fit Score
        # and a live Priority Score formula — strategic value / ease default
        # to 0 rather than blocking the formula with "unknown".
        ws = self.wb["Grant Opportunities"]
        self.assertEqual(ws["A4"].value, "Example Discovery-Only Programme")
        self.assertEqual(ws["D4"].value, 4)
        self.assertEqual(ws["F4"].value, 0)
        self.assertEqual(ws["H4"].value, 0)
        self.assertEqual(ws["B4"].value, "=(2*D4)+F4+H4")

    def _sheet_text(self, name: str) -> str:
        return "\n".join(
            str(cell.value)
            for row in self.wb[name].iter_rows()
            for cell in row
            if cell.value is not None
        )


if __name__ == "__main__":
    unittest.main()
